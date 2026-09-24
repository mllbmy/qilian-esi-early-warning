# -*- coding: utf-8 -*-
"""JMS 版自检：版式 / 匿名 / 关键词 / 经纬度 / 文献 / 图表顺序 / 数值锚点 / 与 EN 稿逐段对比"""
import io
import os
import re
import sys

from docx import Document

# 路径自解析：仓库克隆/解压到任意目录均可运行
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDIR = os.path.join(BASE, "out", "jms")
MAIN = os.path.join(OUTDIR, "论文_EN_JMSv3_2026-09-23.docx")
TP = os.path.join(OUTDIR, "Title_Page_JMSv3_2026-09-23.docx")
TXT = os.path.join(OUTDIR, "论文_EN_JMSv3_2026-09-23.txt")

FAIL, OK = [], []


def chk(cond, msg, detail=""):
    (OK if cond else FAIL).append(msg + (" | " + detail if detail else ""))
    print(("  [OK] " if cond else "  [FAIL] ") + msg + ("  " + detail if detail else ""))


d = Document(MAIN)
paras = [p.text for p in d.paragraphs]
body = "\n".join(paras)
txt = io.open(TXT, encoding="utf-8").read()

print("== 1. 版式（JMS：A4 / 10 pt / 单倍行距 / 页码）==")
s = d.sections[0]
chk(abs(s.page_width.cm - 21.0) < 0.2 and abs(s.page_height.cm - 29.7) < 0.2,
    "A4 页面", "%.1f x %.1f cm" % (s.page_width.cm, s.page_height.cm))
chk(d.styles["Normal"].font.size is not None and abs(d.styles["Normal"].font.size.pt - 10) < 0.01,
    "正文字号 10 pt", str(d.styles["Normal"].font.size.pt))
chk(abs(d.styles["Normal"].paragraph_format.line_spacing - 1.0) < 0.01,
    "单倍行距", str(d.styles["Normal"].paragraph_format.line_spacing))
ftr = s.footer.paragraphs[0]._p.xml if s.footer.paragraphs else ""
chk("PAGE" in ftr, "页脚自动页码域")

print("== 2. 匿名（主文档不得含作者/单位/邮箱）==")
# 禁用词表来自本地私有模块 author_local.py（公开库不含作者身份信息）
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
try:
    from author_local import FORBIDDEN_AUTHOR_TOKENS as _BAD_TOKENS
except ImportError:
    _BAD_TOKENS = ()
if _BAD_TOKENS:
    for bad in _BAD_TOKENS:
        chk(bad not in body, "主文档不含 '%s'" % bad)
else:
    chk("@" not in body, "主文档不含邮箱（未提供本地禁用词表，降级为通用检查）")

print("== 3. 关键词 / 经纬度 / 声明 ==")
kw = [p for p in paras if p.startswith("Keywords:")]
chk(len(kw) == 1, "关键词行唯一")
if kw:
    items = [x.strip() for x in kw[0].replace("Keywords:", "").split(";") if x.strip()]
    chk(len(items) == 6, "关键词 6 个", " / ".join(items))
    chk("arid" not in kw[0].lower(), "已删 arid--semiarid region")
chk("93.5" in body and "40.5" in body, "研究区经纬度已写入正文")
for h in ("Author contribution", "Data availability", "Competing interests", "Ethical statements"):
    chk(h in body, "声明小节：%s" % h)
chk("Conflict of Interest: The author declares that he has no conflict of interest." in body,
    "竞争利益官方模板句")
chk("did not involve human participants or animals" in body, "伦理声明措辞")

print("== 4. 参考文献（19 条 / JMS 体例 / 位置）==")
refs_jms = [l.strip() for l in io.open(os.path.join(BASE, "jms", "refs_jms.txt"), encoding="utf-8").read().split("\n") if l.strip()]
found = [r for r in refs_jms if r in body]
chk(len(found) == 19, "19 条 JMS 体例文献逐字落入主文档", "%d/19" % len(found))
chk(body.count(" doi:") == 0, "无 elsarticle 式 'doi:' 残留")
chk("(In Chinese)" not in body, "文献体例：无 '(In Chinese)' 误标（本论文 19 条均为英文文献）")
i_ref = next((i for i, p in enumerate(paras) if p.strip() == "References"), -1)
i_fig = next((i for i, p in enumerate(paras) if p.strip() == "Figures"), -1)
i_tab = next((i for i, p in enumerate(paras) if p.strip() == "Tables"), -1)
chk(i_ref > 0 and i_fig > i_ref and i_tab > i_fig, "顺序：References → Figures → Tables",
    "ref=%d fig=%d tab=%d" % (i_ref, i_fig, i_tab))
_data_tables = [t for t in d.tables if len(t.rows) >= 2]
chk(len(_data_tables) == 8, "数据表 8 个（Word 可编辑，非图片）",
    "%d 张数据表 + %d 个公式容器 = %d" % (len(_data_tables), len(d.tables) - len(_data_tables), len(d.tables)))

print("== 5. 正文引注体例（JMS：(Author Year)，年份前无逗号）==")
bad_cite = re.findall(r"\(([A-Z][A-Za-z\-]+(?:\s+(?:et al\.|and)\s+[A-Za-z\-]+)?),\s+(?:19|20)\d\d", body)
chk(not bad_cite, "无 'Author, Year' 旧体例", "残留: %s" % sorted(set(bad_cite))[:6])
good_cite = re.findall(r"\(([A-Z][A-Za-z\-]+(?:\s+et\s+al\.)?\s+(?:19|20)\d\d)", body)
chk(len(good_cite) >= 20, "新体例引注数量", str(len(good_cite)))

print("== 6. 数值锚点（红线：一个都不许变）==")
ANCHORS = ["0.052", "0.035", "0.597", "0.789", "0.038", "0.029", "0.44", "0.89",
           "0.098", "0.067", "1.086", "0.632", "0.446", "0.394", "0.510", "0.493",
           "0.488", "0.358", "0.533", "0.306", "0.482", "0.486", "0.709", "0.258",
           "0.643", "0.005", "0.80", "0.71", "0.51", "0.17", "0.849", "0.833",
           "0.818", "0.808", "0.723", "0.529", "0.263", "1.120", "0.0032", "0.034",
           "0.065", "0.425", "0.578", "0.0034", "5.015", "0.0038", "1.12",
           "1914", "999", "0.364", "0.233", "0.221", "0.151", "0.206",
           "0.245", "0.171", "0.136", "0.128"]
_full = body + "\n" + "\n".join(
    c.text for t in d.tables for row in t.rows for c in row.cells)
# 归一化千分位与细空格：1,914 / 1 914 → 1914（只处理 3 位数字组）
_full_norm = re.sub(r"(?<=\d)[,\s\u2009\u00a0](?=\d\d\d(?!\d))", "", _full)
missing = [a for a in ANCHORS if a not in _full_norm]
chk(not missing, "全部 %d 个数值锚点在位" % len(ANCHORS), "缺失: %s" % missing)

print("== 7. 与既有 EN 稿逐段对比（应只差：引注逗号 / 经纬度句 / 关键词 / 去题头）==")
OLD = os.path.join(BASE, "out", "论文_EN_review.docx")
if os.path.isfile(OLD):
    od = Document(OLD)
    op = [p.text.strip() for p in od.paragraphs if p.text.strip()]
    np_ = [p.strip() for p in paras if p.strip()]
    def _norm(x):
        x = re.sub(r"\(\s*(?:[A-Z][^()]{0,60}?(?:19|20)\d\d[^()]{0,40})?\)", " ", x)
        x = re.sub(r"[A-Z][A-Za-z\-]+(?:\s+et\s+al\.)?(?:,\s*|\s+)(?:19|20)\d\d[a-z]?", "CIT", x)
        x = re.sub(r"\s+", " ", x)
        return x.strip()

    op_n = {_norm(p) for p in op}
    np_n = {_norm(p) for p in np_}
    only_old = [p for p in op if _norm(p) not in np_n]
    print("    旧稿段落 %d / 新稿段落 %d / 旧稿独有 %d 段" % (len(op), len(np_), len(only_old)))
    for p in only_old[:6]:
        print("      -", p[:95])
else:
    print("    （旧 EN Word 不在，跳过对比）")

print()
print("结论: %d 项通过, %d 项失败" % (len(OK), len(FAIL)))
for f in FAIL:
    print("  ! ", f)
sys.exit(1 if FAIL else 0)