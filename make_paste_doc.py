# -*- coding: utf-8 -*-
r"""生成《论文中文段落_可直接粘贴.docx》——段内公式为**真正的 Word 公式对象（OMML）**，
点击即进入公式编辑器（与投稿主文档中的式 (1) 同源，复用 word_style 的 OMML 组件）。

用法：python make_paste_doc.py
输出：交付\JMS_v5_2026-09-28\论文中文段落_可直接粘贴.docx
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn

from word_style import (OMATH, R, RU, SUB, SUP, SSUBSUP, FRAC, NARY, SUM, D,
                        EQ3, EQNORM, EQ4, EQ5, add_equation)
from fix_inline_math import fix_inline_math

OUT = r"D:\比赛\论文\交付\JMS_v5_2026-09-28\论文中文段落_可直接粘贴.docx"


def RT(s):
    """数学文本 run：转义 XML 保留字符（<、>、&）。"""
    return R(s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


# ── 行内公式（OMML 主体，与论文记号一致）──
def eq_ej():                       # e_j = −(1/ln n) Σ_i p_ij ln p_ij
    pij = SUB(R("p"), R("ij"))
    return (SUB(R("e"), R("j")) + R(" = −") + FRAC(R("1"), RU("ln ") + R("n"))
            + SUM(R("i"), pij + RU(" ln ") + pij))

def eq_pij():                      # p_ij = x′_ij / Σ_i x′_ij
    xpij = SUB(SUP(R("x"), R("′")), R("ij"))
    return SUB(R("p"), R("ij")) + R(" = ") + FRAC(xpij, SUM(R("i"), xpij))

def eq_wEj():                      # w^E_j = (1 − e_j) / Σ_j (1 − e_j)
    one_ej = D(R("1 − ") + SUB(R("e"), R("j")))
    return (SUB(SUP(R("w"), R("E")), R("j")) + R(" = ")
            + FRAC(one_ej, SUM(R("j"), one_ej)))

def eq_wAj():                      # w^A_j
    return SUB(SUP(R("w"), R("A")), R("j"))

def eq_lmax():                     # λ_max = 5.015
    return SUB(R("λ"), RU("max")) + R(" = 5.015")

def eq_ci():                       # CI = (λ_max − n)/(n − 1) = 0.0038
    lm, n = SUB(R("λ"), RU("max")), R("n")
    return R("CI = ") + FRAC(D(lm + R(" − ") + n), D(n + R(" − 1"))) + R(" = 0.0038")

def eq_cr():                       # CR = CI/RI = 0.0034 < 0.1
    return R("CR = ") + FRAC(R("CI"), R("RI")) + RT(" = 0.0034 < 0.1")

TITLE = "论文中文段落（可直接复制粘贴）"
NOTE = ("段内公式为真正的 Word 公式对象：点击即进入公式编辑器，可像正常公式一样修改；"
        "整段复制粘贴到目标文档后仍保持公式对象。")

P_CAP = ["图1  祁连山区域地形、祁连山国家公园及县域位置",
         "图幅范围即分析窗口 36.0-40.5°N，93.5-103.5°E，红点为 11 个县级分析单元；虚线轮廓为据《祁连山国家公园总体规划》"
         "数字化的公园范围（审图号 GS(2019)46 号，仅示意；数字化面积 48,800 km²，与规划总面积 50,200 km² 相差 3% 以内）；"
         "底图为 Copernicus DEM 90 m，省界来自 Natural Earth 1:50m（仅作参考）。"]
P_DATA = "本研究的数据源汇总于表 1。土地覆被分 2005、2010、2015、2020、2023 五期计算，并线性插值到逐年值。" \
         "气候格网分别聚合到区域尺度与 11 个县域质心，故县域数值为质心窗口估计，而非全县面积加权统计。"
P_DPSIR = ("遵循 DPSIR 框架 (Hou et al., 2024; Yang et al., 2023)，九项指标覆盖各准则层：驱动力（降水、气温）、"
           "压力（裸地/退化占比、耕地占比）、状态（草地、森林、水域占比）、影响（草地变化率）、响应（生态用地占比）。"
           "指标选取兼顾两点：仅用公开可复现数据（CRU TS 4.10 气候与 CLCD 30 m 土地覆被），"
           "并同时刻画干旱—半干旱高寒生态系统中主导生态安全的气候驱动与地表响应。"
           "气温与降水均取正向：在寒冷干旱的高寒环境中，热量与水分同为植被生长的限制因子。"
           "升温带来的蒸散发增强等相反效应，则在空间归因分析中讨论。"
           "表 2 列出各指标的方向与角色，图 2 给出 DPSIR 因果链与 ESI 预警管线的衔接；赋权前所有指标均做 min-max 归一化到 [0,1]。")
SPARE = ("纯文本备用（公式为平字）：熵权为 ej = −(1/ln n) Σi pij ln pij，其中 pij = x′ij / Σi x′ij，"
         "wEj = (1 − ej) / Σj(1 − ej)。AHP 权重 wAj 采用主特征向量法 (Saaty, 1980)。"
         "因生态安全最直接由状态层与响应层体现，准则层判断矩阵（5 个准则，Saaty 1-9 标度；表 3）"
         "依 DPSIR 层次优先级（状态 > 响应 > 影响 > 驱动力 = 压力）构造，"
         "得 λmax = 5.015、CI = (λmax − n)/(n − 1) = 0.0038、CR = CI/RI = 0.0034 < 0.1（n = 5 时 RI = 1.12），"
         "满足一致性要求。各准则层内指标均分该层权重。组合权重为：")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

h = doc.add_heading(TITLE, level=1)
p = doc.add_paragraph()
r = p.add_run(NOTE)
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)


def _add_eq(par, body):
    par._p.append(OMATH(body))


# ① 图 1 题注
doc.add_heading("图 1 题注", level=2)
for s in P_CAP:
    doc.add_paragraph(s)

# ② 2.2 数据来源
doc.add_heading("2.2 数据来源（首段）", level=2)
doc.add_paragraph(P_DATA)

# ③ 3.1 DPSIR
doc.add_heading("3.1 DPSIR 框架与指标体系（首段）", level=2)
doc.add_paragraph(P_DPSIR)

# ④ 3.1 熵权与 AHP 权重段（公式为公式对象）
doc.add_heading("3.1 熵权与 AHP 权重段", level=2)
par = doc.add_paragraph()
par.add_run("熵权为 ")
_add_eq(par, eq_ej())
par.add_run("，其中 ")
_add_eq(par, eq_pij())
par.add_run("，")
_add_eq(par, eq_wEj())
par.add_run("。AHP 权重 ")
_add_eq(par, eq_wAj())
par.add_run(" 采用主特征向量法 (Saaty, 1980)。因生态安全最直接由状态层与响应层体现，"
            "准则层判断矩阵（5 个准则，Saaty 1-9 标度；表 3）依 DPSIR 层次优先级"
            "（状态 > 响应 > 影响 > 驱动力 = 压力）构造，得 ")
_add_eq(par, eq_lmax())
par.add_run("、")
_add_eq(par, eq_ci())
par.add_run("、")
_add_eq(par, eq_cr())
par.add_run("（n = 5 时 RI = 1.12），满足一致性要求。各准则层内指标均分该层权重。组合权重为：")

# ⑤ 式 (1)：组合权重（与投稿稿同源，带编号）
add_equation(doc, EQ3(), 1)

sp = doc.add_paragraph()
rr = sp.add_run(SPARE)
rr.font.size = Pt(8.5)
rr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

# ⑥ 3.3 生态安全指数（ESI）—— 式 (2)(3) 为公式对象，编号与投稿稿一致
doc.add_heading("3.3 生态安全指数（ESI）", level=2)
doc.add_paragraph("指标先按方向做 min-max 标准化到 [0,1]（式 2）：正向指标值越大越安全，负向指标反之。")
add_equation(doc, EQNORM(), 2)
doc.add_paragraph("ESI 为各指标标准化值的加权和：")
add_equation(doc, EQ4(), 3)

# ⑦ 3.4 预警预测模型与等级划分 —— 式 (4) 为公式对象
doc.add_heading("3.4 预警预测模型与等级划分", level=2)
doc.add_paragraph("GM(1,1) 模型 (Deng, 1982) 以一次累加序列为基础，用最小二乘估计参数（â, b̂）：")
add_equation(doc, EQ5(), 4)
doc.add_paragraph(
    "GM(1,1) 假定单调指数趋势，而观测 ESI 序列在 2010 年见顶后回落，违背该假定；"
    "本文因此额外采用以幂指数 γ 推广 GM(1,1) 的非线性灰色伯努利模型（NGBM），并与 ARIMA 对比。"
    "模型精度由后验差比值 C=S2/S1 与小误差概率 P 评价：优（C<0.35，P>0.95）、"
    "合格（C<0.50，P>0.80）、勉强合格（C<0.65，P>0.70）。"
    "模型精度另由滚动回测验证：以 2005-2018 年拟合，用留出的 2019-2023 年观测评估。")
doc.add_paragraph(
    "预警等级分为五级：安全 [0.8,1]、较安全 [0.6,0.8)、轻警 [0.4,0.6)、中警 [0.2,0.4)、重警 [0,0.2)。"
    "等级界限在归一化区间 [0,1] 上按 0.2 等距划分，与 DPSIR 类保护区生态安全预警框架的五级划分惯例一致 "
    "(Liu et al., 2023)；等距划分便于跨研究比较，且避免样本相关的断点。"
    "后文将表明，观测序列的预警等级对权重比例不敏感。")

# ⑧ 4.1 指标原始动态（图 3 段）
doc.add_heading("4.1 指标原始动态（图 3 段）", level=2)
doc.add_paragraph(
    "图 3 展示 ESI 背后的指标原始动态。降水量在 166.9-284.4 mm/年之间波动，"
    "由 2005 年的 228.0 mm 降至 2023 年的 166.9 mm；年均气温由 1.29°C 升至 1.92°C。"
    "土地覆被类指标（在五个 CLCD 时点间线性插值）变化缓慢：生态用地占比（草地+森林+水域）"
    "由 2005 年的 37.9% 升至 2010 年的 39.6%，2023 年仍接近 39.0%；"
    "裸地/退化占比先由 58.2% 降至 56.5%，随后微升至 57.0%；"
    "水域（1.75%-2.12%）与森林（1.15%-1.29%）占比稳步上升，耕地小幅增加（3.8%-4.0%）。"
    "草地变化率波动最大（-0.47% 至 +0.84%/年）。"
    "这些序列经标准化与组合赋权（3.3 节）后，构成下文分析的 ESI 轨迹（图 4）。")

# ⑨ 图 3 题注
doc.add_heading("图 3 题注", level=2)
doc.add_paragraph(
    "图3  九项 DPSIR 指标变化趋势（2005-2023；圆点为五个 CLCD 土地覆被观测时点 "
    "2005、2010、2015、2020、2023，土地覆被类指标在时点间线性插值；虚线为 2010 年即 ESI 峰值年）")

# ⑩ 4.1 指标权重
doc.add_heading("4.1 指标权重", level=2)
doc.add_paragraph(
    "熵权结果（表 4）显示，草地变化率（w=0.245）与耕地占比（w=0.171）的信息量最大。"
    "AHP 准则层权重突出状态层指标（草地、森林、水域，各 0.146）与生态用地占比（0.250）；"
    "组合权重（α=0.4）下，草地变化率（0.206）、生态用地占比（0.136）、水域占比（0.128）"
    "与森林占比（0.120）最高，表明草地动态、生态用地规模与水资源共同主导生态安全的时序变化（图 7）。")

n_fix = fix_inline_math(doc)
doc.save(OUT)
print("  ✓ %s（行内数学重排 %d 处 run）" % (OUT, n_fix))
print("  ✓ 大小 %.1f KB" % (os.path.getsize(OUT) / 1024))