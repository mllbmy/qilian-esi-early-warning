# -*- coding: utf-8 -*-
"""行内公式排版器：把 Word 稿件里被压平成纯文本的行内数学符号
（ej / pij / x'ij / wEj / wAj / Σi / Σj / λmax / CI / CR / RI / S2 / S1 / α / γ / q）
重新排成 斜体 + 上标 + 下标，符合期刊"行内符号用斜体"的通行排版惯例。

用法：
    from fix_inline_math import fix_inline_math
    fix_inline_math(doc)        # 就地修改 python-docx Document，随后自行 save
或命令行：
    python fix_inline_math.py <file.docx>
"""
import copy
import re
import sys

from docx.oxml.ns import qn

# ── 令牌表：(匹配文本, [(文本, 斜体, 下标, 上标), ...])，长令牌在前 ──
_I = 1
TOKENS = [
    ("x'ij", [("x", _I, 0, 0), ("\u2032", 0, 0, 0), ("ij", 0, 1, 0)]),
    ("wEj",  [("w", _I, 0, 0), ("E", _I, 0, 1), ("j", _I, 1, 0)]),
    ("wAj",  [("w", _I, 0, 0), ("A", _I, 0, 1), ("j", _I, 1, 0)]),
    ("pij",  [("p", _I, 0, 0), ("ij", 0, 1, 0)]),
    ("\u03a3i", [("\u03a3", 0, 0, 0), ("i", _I, 1, 0)]),
    ("\u03a3j", [("\u03a3", 0, 0, 0), ("j", _I, 1, 0)]),
    ("\u03bbmax", [("\u03bb", _I, 0, 0), ("max", 0, 1, 0)]),
    ("wE",   [("w", _I, 0, 0), ("E", _I, 0, 1)]),
    ("wA",   [("w", _I, 0, 0), ("A", _I, 0, 1)]),
    ("ej",   [("e", _I, 0, 0), ("j", _I, 1, 0)]),
    ("S2",   [("S", _I, 0, 0), ("2", 0, 1, 0)]),
    ("S1",   [("S", _I, 0, 0), ("1", 0, 1, 0)]),
    ("CI",   [("CI", _I, 0, 0)]),
    ("CR",   [("CR", _I, 0, 0)]),
    ("RI",   [("RI", _I, 0, 0)]),
]
# ── 正则令牌（需上下文约束，避免误伤普通单词）──
REGEX_TOKENS = [
    (re.compile(r"(?<![A-Za-z])q(?==)"), [("q", _I, 0, 0)]),
    (re.compile(r"(?<=ln )n(?![A-Za-z])"), [("n", _I, 0, 0)]),
    (re.compile(r"\u03b1"), [("\u03b1", _I, 0, 0)]),
    (re.compile(r"\u03b3"), [("\u03b3", _I, 0, 0)]),
]

_SPLIT = re.compile("(" + "|".join(re.escape(k) for k, _ in TOKENS) + ")")
_SKIP_STYLE = {"Heading 1", "Heading 2", "Heading 3", "Title"}


def _spec(text):
    """把一段文本切成 [(文本, 斜体, 下标, 上标), ...]。"""
    out = []
    for piece in _SPLIT.split(text):
        if not piece:
            continue
        mapped = dict(TOKENS).get(piece)
        if mapped is not None:
            out.extend(mapped)
            continue
        # 正则令牌：在同一段内二次切分
        cuts, last = [], 0
        for rx, rep in REGEX_TOKENS:
            for m in rx.finditer(piece):
                if m.start() >= last:
                    cuts.append((m.start(), m.end(), rep))
                    last = m.end()
        if not cuts:
            out.append((piece, None, None, None))
            continue
        pos = 0
        for s, e, rep in cuts:
            if s > pos:
                out.append((piece[pos:s], None, None, None))
            out.extend(rep)
            pos = e
        if pos < len(piece):
            out.append((piece[pos:], None, None, None))
    return out


def _format_paragraph(par):
    """就地把段落内 run 的行内数学符号重排；返回改动次数。"""
    if par.style is not None and par.style.name in _SKIP_STYLE:
        return 0
    changed = 0
    for run in list(par.runs):
        text = run.text
        if not text or not (_SPLIT.search(text) or any(rx.search(text) for rx, _ in REGEX_TOKENS)):
            continue
        parts = _spec(text)
        if all(p[1] is None and p[2] is None and p[3] is None for p in parts):
            continue
        src = run._r
        base = copy.deepcopy(src)
        for child in list(base):
            if child.tag == qn("w:t"):
                base.remove(child)
        anchor = src
        new_runs = []
        for txt, ital, sub, sup in parts:
            r = copy.deepcopy(base)
            t = r.makeelement(qn("w:t"), {})
            t.text = txt
            t.set(qn("xml:space"), "preserve")
            r.append(t)
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                rpr = r.makeelement(qn("w:rPr"), {})
                r.insert(0, rpr)
            for tag in ("w:i", "w:iCs", "w:vertAlign"):
                for el in rpr.findall(qn(tag)):
                    rpr.remove(el)
            if ital:
                for tag in ("w:i", "w:iCs"):
                    el = rpr.makeelement(qn(tag), {})
                    rpr.append(el)
            if sub or sup:
                va = rpr.makeelement(qn("w:vertAlign"), {qn("w:val"): "subscript" if sub else "superscript"})
                rpr.append(va)
            anchor.addnext(r)
            anchor = r
            new_runs.append(r)
        src.getparent().remove(src)
        changed += 1
    return changed


def fix_inline_math(doc):
    """处理正文段落与所有表格单元格，返回改动 run 数。"""
    total = 0
    for par in doc.paragraphs:
        total += _format_paragraph(par)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    total += _format_paragraph(par)
    return total


if __name__ == "__main__":
    if sys.version_info[0] < 3:
        raise SystemExit("python3 required")
    from docx import Document

    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    for path in sys.argv[1:]:
        d = Document(path)
        n = fix_inline_math(d)
        d.save(path)
        print("  %s：重排 %d 处 run" % (path, n))