# -*- coding: utf-8 -*-
"""Build the current anonymous submission document from manuscript.tex and .bbl."""
import io
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

try:
    import pymupdf
except ImportError:
    pymupdf = None
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import docxbuild_lib as L
import omml

STATS = {"unmapped_cmd": {}, "warnings": [], "figures": [], "tables": [],
         "equations": [], "images_missing": []}

# --------------------------------------------------------------------------- #
# document setup
# --------------------------------------------------------------------------- #
def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = s.bottom_margin = Cm(2.5)
    add_page_numbers(s)
    n = doc.styles["Normal"]
    n.font.name = "Times New Roman"
    n.font.size = Pt(L.BODY_PT)
    rpr = n.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Times New Roman")
    pf = n.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    return doc


def add_page_numbers(section):
    """Automatic page number in the footer (JMS guide p8). The PAGE field is the
    only field function in the document."""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.font.size = Pt(L.BODY_PT)
    set_run_fonts(run)
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld); run._r.append(instr); run._r.append(fld2)


def set_run_fonts(run, name="Times New Roman"):
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)


def para(doc, pieces_or_text, size=L.BODY_PT, align=None, bold=False,
         italic=False, before=0, after=0, indent=None, hanging=None,
         keep_with_next=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    if indent is not None:
        pf.left_indent = Cm(indent)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)
    if keep_with_next:
        pf.keep_with_next = True
    if isinstance(pieces_or_text, str):
        r = p.add_run(pieces_or_text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        set_run_fonts(r)
    else:
        L.emit_pieces(p, pieces_or_text, size=size)
    return p


# --------------------------------------------------------------------------- #
# figures
# --------------------------------------------------------------------------- #
def rasterise(pdf_path, dpi=300):
    os.makedirs(L.IMGDIR, exist_ok=True)
    stem = os.path.splitext(os.path.basename(pdf_path))[0]
    out = os.path.join(L.IMGDIR, "%s_%ddpi.png" % (stem, dpi))
    if os.path.exists(out) and os.path.getmtime(out) > os.path.getmtime(pdf_path):
        return out
    if pymupdf is None:
        import subprocess
        subprocess.run([os.environ.get("QILIAN_PDF_PYTHON", r"D:\Program Files\Python313\python.exe"),
                        os.path.join(os.path.dirname(__file__), "pdf_support.py"),
                        pdf_path, out, str(dpi)], check=True, capture_output=True)
        return out
    d = pymupdf.open(pdf_path)
    pm = d[0].get_pixmap(matrix=pymupdf.Matrix(dpi / 72.0, dpi / 72.0))
    pm.save(out)
    d.close()
    # keep the embedded picture a sane size
    while os.path.getsize(out) > 4_500_000 and dpi > 170:
        dpi = int(dpi * 0.8)
        pm = pymupdf.open(pdf_path)[0].get_pixmap(
            matrix=pymupdf.Matrix(dpi / 72.0, dpi / 72.0))
        pm.save(out)
    return out


def emit_figure(doc, block, biblio, numbering):
    inc = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]*)\}", block)
    cap = re.search(r"\\caption\{", block)
    label = L._label_of(block)
    num = numbering.label.get(label, str(numbering.n_fig))
    if not inc:
        STATS["images_missing"].append("(no \\includegraphics) " + str(label))
        return
    rel = inc.group(1)
    src = os.path.join(L.PKG, rel)
    if not os.path.exists(src):
        STATS["images_missing"].append(rel)
        para(doc, "[missing figure file: %s]" % rel, bold=True)
    else:
        png = rasterise(src)
        from PIL import Image
        with Image.open(png) as im:
            aspect = im.width / im.height
        w = L.TEXTW
        h = w / aspect
        if h > 21.0:
            h = 21.0
            w = h * aspect
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(png, width=Cm(w))
        STATS["figures"].append({"n": num, "label": label, "file": rel,
                                 "png": os.path.basename(png),
                                 "w_cm": round(w, 2), "h_cm": round(h, 2)})
    # caption
    if cap:
        inner, _ = L.find_group(block, cap.end() - 1)
        pieces = L.resolve_refs(L.adopt_bases(L.parse_inline(inner, stats=STATS["unmapped_cmd"])),
                                biblio, numbering, STATS["warnings"])
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.keep_together = True
        r = p.add_run("Figure %s: " % num)
        r.bold = True; r.font.size = Pt(L.BODY_PT); set_run_fonts(r)
        L.emit_pieces(p, pieces, size=L.BODY_PT)


# --------------------------------------------------------------------------- #
# tables
# --------------------------------------------------------------------------- #
W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _tbl_borders(table, top_sz=12, bot_sz=12):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    xml = ('<w:tblBorders %s>'
           '<w:top w:val="single" w:sz="%d" w:space="0" w:color="000000"/>'
           '<w:bottom w:val="single" w:sz="%d" w:space="0" w:color="000000"/>'
           '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
           '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
           '<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
           '<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
           '</w:tblBorders>' % (W_NS, top_sz, bot_sz))
    tblPr.append(parse_xml(xml))


def cell_bottom(cell, sz=6, val="single"):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    tcPr.append(parse_xml(
        '<w:tcBorders %s><w:bottom w:val="%s" w:sz="%d" w:space="0" '
        'w:color="000000"/></w:tcBorders>' % (W_NS, val, sz)))


def set_col_widths(table, widths_cm):
    tbl = table._tbl
    tblPr = tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
            gc.set(qn("w:w"), str(int(w * 567)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def emit_table(doc, block, biblio, numbering):
    label = L._label_of(block)
    num = numbering.label.get(label, str(numbering.n_tab))
    cap = re.search(r"\\caption\{", block)
    tabm = re.search(r"\\begin\{tabular\}(?:\{)?", block)
    # caption first
    if cap:
        inner, _ = L.find_group(block, cap.end() - 1)
        pieces = L.resolve_refs(L.adopt_bases(L.parse_inline(inner, stats=STATS["unmapped_cmd"])),
                                biblio, numbering, STATS["warnings"])
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
        r = p.add_run("Table %s: " % num)
        r.bold = True; r.font.size = Pt(L.BODY_PT); set_run_fonts(r)
        L.emit_pieces(p, pieces, size=L.BODY_PT)
    # tabular body
    spec = ""
    j = block.find("\\begin{tabular}")
    k = j + len("\\begin{tabular}")
    if k < len(block) and block[k] == "{":
        spec, body_start = L.find_group(block, k)
    else:
        body_start = k
    end = block.find("\\end{tabular}")
    body = block[body_start:end]
    rows, ncol = L.parse_tabular(spec, body)
    cols = L.parse_colspec(spec)
    # a source table may declare more columns than any row actually uses
    # (tab:county declares lcccc but every row has 4 cells): render the used ones
    max_span = 1
    for cells, _rules in rows:
        s = 0
        for c in cells:
            mm = re.match(r"\\multicolumn\{(\d+)\}", c.strip())
            s += int(mm.group(1)) if mm else 1
        max_span = max(max_span, s)
    if max_span != ncol:
        STATS["warnings"].append(
            "table %s: colspec declares %d columns, rows use %d -> using %d"
            % (num, ncol, max_span, max_span))
        ncol = max_span
    if len(cols) != ncol:
        STATS["warnings"].append("table %s colspec/col mismatch %s" % (num, spec))
        cols = [(cols[k][0] if k < len(cols) else "c", None) for k in range(ncol)]
    widths = L.column_widths(cols, ncol)
    scale = L.TEXTW / sum(widths) if sum(widths) > L.TEXTW else 1.0
    widths = [w * scale for w in widths]

    parsed_rows = []
    for cells, rules in rows:
        parsed_rows.append(([parse_cell(c) for c in cells], rules))
    header_rows = 1
    for idx, (_, rules) in enumerate(parsed_rows):
        if any(r[0] == "mid" for r in rules):
            header_rows = idx + 1
            break
    if header_rows == 1 and len(parsed_rows) > 1:
        header_rows = 1

    table = doc.add_table(rows=len(parsed_rows), cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _tbl_borders(table)
    set_col_widths(table, widths)

    for ri, (cells, rules) in enumerate(parsed_rows):
        row = table.rows[ri]
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if ri < header_rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        col = 0
        for txt, span, align in cells:
            if col >= ncol:
                break
            if span > 1:
                span = min(span, ncol - col)
                merged = row.cells[col].merge(row.cells[col + span - 1])
            else:
                merged = row.cells[col]
            target = merged
            target.paragraphs[0].paragraph_format.space_after = Pt(0)
            target.paragraphs[0].paragraph_format.line_spacing = 1.0
            # Keep short tables intact and long-table headers with their first row.
            target.paragraphs[0].paragraph_format.keep_with_next = (
                ri < header_rows or (len(parsed_rows) <= 14 and ri < len(parsed_rows)-1))
            if align:
                target.paragraphs[0].alignment = align
            pieces = L.resolve_refs(
                L.adopt_bases(L.parse_inline(txt, stats=STATS["unmapped_cmd"])),
                biblio, numbering, STATS["warnings"])
            if ri < header_rows:
                pieces = [(p[0], p[1], dict(p[2], b=True)) if p[0] == "text" else p
                          for p in pieces]
            L.emit_pieces(target.paragraphs[0], pieces, size=L.TABLE_PT)
            col += span
        # rules under this row
        for kind, arg in rules:
            if kind == "mid":
                for c in row.cells:
                    cell_bottom(c, sz=6)
            elif kind == "cmid" and arg:
                try:
                    a, b = arg.split("-")
                    for ci in range(int(a) - 1, int(b)):
                        if ci < ncol:
                            cell_bottom(row.cells[ci], sz=4)
                except Exception:
                    pass
    # notes after \end{tabular}
    tail = block[end + len("\\end{tabular}"):]
    tail = tail.replace("\\end{table}", "")
    tail = re.sub(r"\\vspace\{[^}]*\}", " ", tail)
    tail_pieces = L.adopt_bases(L.parse_inline(tail, stats=STATS["unmapped_cmd"]))
    tail_text = "".join(p[1] for p in tail_pieces if p[0] == "text").strip()
    if tail_text:
        for cell in table.rows[-1].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
        pieces = L.resolve_refs(tail_pieces, biblio, numbering, STATS["warnings"])
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.keep_together = True
        L.emit_pieces(p, pieces, size=L.NOTE_PT)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(10)
    STATS["tables"].append({"n": num, "label": label, "cols": ncol,
                            "rows": len(parsed_rows), "header_rows": header_rows,
                            "spec": spec,
                            "widths": [round(w, 2) for w in widths]})


def parse_cell(txt):
    """Return list of (text, span, align) for one tabular cell."""
    txt = txt.strip()
    m = re.match(r"\\multicolumn\{(\d+)\}\{([^}]*)\}\{", txt)
    if m:
        span = int(m.group(1))
        align_spec = m.group(2)
        inner, _ = L.find_group(txt, m.end() - 1)
        align = {"c": WD_ALIGN_PARAGRAPH.CENTER,
                 "r": WD_ALIGN_PARAGRAPH.RIGHT,
                 "l": WD_ALIGN_PARAGRAPH.LEFT}.get(
                     align_spec.strip()[:1], WD_ALIGN_PARAGRAPH.CENTER)
        return (inner, span, align)
    return (txt, 1, None)


# --------------------------------------------------------------------------- #
# equations
# --------------------------------------------------------------------------- #
def emit_equation(doc, latex, num, biblio, numbering):
    latex = re.sub(r"\\label\{[^}]*\}", "", latex)
    latex = " ".join(latex.split())
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.tab_stops.add_tab_stop(Cm(L.TEXTW / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(L.TEXTW), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    p._p.append(omml.omml_element(latex, display=True))
    if num is not None:
        p.add_run("\t(%d)" % num)
    STATS["equations"].append({"n": num, "latex": latex,
                               "text": omml.omml_text(latex, display=True)})


# --------------------------------------------------------------------------- #
# main build
# --------------------------------------------------------------------------- #
def build():
    tex = L.strip_comments(L.read(L.TEX))
    bbl = L.read(L.BBL)
    biblio = L.Bib(bbl)

    fm_start = tex.find("\\begin{frontmatter}")
    title = re.search(r"\\title\{(.*?)\}\s*\n", tex, re.S).group(1)
    title = L.norm_ws(title)
    abs_m = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", tex, re.S)
    abstract = L.norm_ws(abs_m.group(1))
    kw_m = re.search(r"\\begin\{keyword\}(.*?)\\end\{keyword\}", tex, re.S)
    keywords = [k.strip() for k in re.split(r"\\sep", L.norm_ws(kw_m.group(1))) if k.strip()]

    body = tex[tex.find("\\end{frontmatter}") + len("\\end{frontmatter}"):]
    body = body.split("\\end{document}")[0]
    numbering = L.assign_numbers(body)

    doc = new_doc()

    # ---- title block (anonymous: JMS double-blind review) ----
    # Author names, affiliation, e-mail and ORCID belong to the separate Title
    # Page file; the main document must not identify the author (guide p12).
    para(doc, title, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, after=12)

    # ---- abstract ----
    para(doc, "Abstract", size=11, bold=True, before=6, after=4)
    ap = L.adopt_bases(L.parse_inline(abstract, stats=STATS["unmapped_cmd"]))
    para(doc, ap, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6)
    kp = []
    for i, k in enumerate(keywords):
        if i:
            kp.append(("text", "; ", L.STYLE0))
        kp.extend(L.adopt_bases(L.parse_inline(k, stats=STATS["unmapped_cmd"])))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Keywords: "); r.bold = True; r.font.size = Pt(L.BODY_PT); set_run_fonts(r)
    L.emit_pieces(p, kp)

    # ---- body ----
    eq_no = 0
    pos = 0
    main_figs, main_tabs = [], []
    in_appendix = False
    for m in L.SPECIAL.finditer(body):
        text = body[pos:m.start()]
        emit_text_chunk(doc, text, biblio, numbering)
        t = m.group(0)
        if t.startswith("\\begin{figure}"):
            if in_appendix:
                emit_figure(doc, t, biblio, numbering)
            else:
                main_figs.append(t)
        elif t.startswith("\\begin{table}"):
            if in_appendix:
                emit_table(doc, t, biblio, numbering)
            else:
                main_tabs.append(t)
        elif t.startswith("\\begin{equation}"):
            inner = t[len("\\begin{equation}"):t.rfind("\\end{equation}")]
            eq_no += 1
            emit_equation(doc, inner, eq_no, biblio, numbering)
        elif t.startswith("\\begin{enumerate}"):
            emit_enumerate(doc, t, biblio, numbering)
        elif t.startswith("\\section") or t.startswith("\\subsection"):
            emit_heading(doc, t, numbering, numbering.by_pos.get(m.start(), ""))
        elif t.startswith("\\appendix"):
            in_appendix = True
        elif t.startswith("\\bibliography"):
            emit_references(doc, biblio)
        pos = m.end()
    emit_text_chunk(doc, body[pos:], biblio, numbering)

    # JMS guide p5: figures (with legends) and tables are placed in sequence at
    # the end of the paper, after the reference list.
    for t in main_figs:
        emit_figure(doc, t, biblio, numbering)
    for t in main_tabs:
        emit_table(doc, t, biblio, numbering)

    if "REPLACE-WITH-MIRROR-ID" in open(L.TEX, encoding="utf-8").read():
        print("!! WARNING: the anonymous-mirror URL is still the placeholder;")
        print("!!          run  python paper-code/set_anon_repo.py <mirror-url>  before submitting")
    L.scrub_metadata(doc, getattr(L, 'TITLE', ''))
    os.makedirs(os.path.dirname(L.OUT), exist_ok=True)
    doc.save(L.OUT)
    return doc, biblio, numbering


def emit_text_chunk(doc, chunk, biblio, numbering):
    chunk = chunk.strip("\n")
    if not chunk.strip():
        return
    for para_txt in re.split(r"\n\s*\n", chunk):
        t = L.norm_ws(para_txt)
        if not t:
            continue
        stripped = re.sub(r"\\(?:label|vspace|setlength|small|footnotesize|"
                          r"scriptsize|centering|noindent|bibliographystyle)"
                          r"(?:\{[^}]*\})?", "", t).strip()
        if not stripped:
            continue
        pieces = L.adopt_bases(L.parse_inline(t, stats=STATS["unmapped_cmd"]))
        pieces = L.resolve_refs(pieces, biblio, numbering, STATS["warnings"])
        if not any(p[0] in ("text", "math", "scr") for p in pieces):
            continue
        p=para(doc, pieces, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4)
        if t.startswith(r'Table~\ref') and len(t)<700:
            p.paragraph_format.keep_with_next = True


def emit_heading(doc, t, numbering, numtxt):
    if t.startswith("\\subsection"):
        level = 2
        body = t[len("\\subsection"):]
    else:
        level = 1
        body = t[len("\\section"):]
    starred = body.startswith("*")
    if starred:
        body = body[1:]
    title = body.strip()
    if title.startswith("{") and title.endswith("}"):
        title = title[1:-1]
    pieces = L.adopt_bases(L.parse_inline(title, stats=STATS["unmapped_cmd"]))
    if numtxt:
        pieces = [("text", "%s  " % numtxt, L.STYLE0)] + pieces
    size = 12 if level == 1 else 11
    para(doc, pieces, size=size, bold=True,
         before=(12 if level == 1 else 8), after=4, keep_with_next=True)


def emit_enumerate(doc, t, biblio, numbering):
    inner = t[t.find("}") + 1:t.rfind("\\end{enumerate}")]
    items = re.split(r"\\item\b", inner)[1:]
    for k, it in enumerate(items, 1):
        pieces = L.adopt_bases(L.parse_inline(L.norm_ws(it), stats=STATS["unmapped_cmd"]))
        pieces = L.resolve_refs(pieces, biblio, numbering, STATS["warnings"])
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run("%d.  " % k); r.font.size = Pt(L.BODY_PT); set_run_fonts(r)
        L.emit_pieces(p, pieces)


def emit_references(doc, biblio):
    para(doc, "References", size=12, bold=True, before=14, after=6, keep_with_next=True)
    for i, (key, label, pieces) in enumerate(biblio.entries, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        L.emit_pieces(p, pieces, size=9)


if __name__ == "__main__":
    doc, biblio, numbering = build()
    print("saved:", L.OUT, os.path.getsize(L.OUT), "bytes")
    print("figures:", len(STATS["figures"]), "tables:", len(STATS["tables"]),
          "equations:", len(STATS["equations"]))
    print("missing images:", STATS["images_missing"])
    print("warnings:", STATS["warnings"][:20], "(%d)" % len(STATS["warnings"]))
    print("unmapped commands:", STATS["unmapped_cmd"])
    print("bib entries:", len(biblio.entries))
    json.dump(STATS, io.open(os.path.join(L.WORK, "build_stats.json"), "w",
                             encoding="utf-8"), indent=1, ensure_ascii=False)
