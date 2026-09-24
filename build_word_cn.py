# -*- coding: utf-8 -*-
"""生成改造后论文的中文版 Word 文档（完整中译，含图、表、公式）"""
import os
import pandas as pd
from word_style import add_equation, make_three_line, heading_cn, EQNORM, EQ3, EQ4, EQ5, EQ6

# 公式编号与顺序严格对齐 EN：式1 组合权重 / 式2 方向标准化(cases) / 式3 ESI / 式4 GM(1,1) / 式5 q
EQMAP = {"eq_weight": EQ3, "eq_norm": EQNORM, "eq_esi": EQ4, "eq_gm": EQ5, "eq_q": EQ6}

def h1(text):
    return heading_cn(doc, text, 15)

def h2(text):
    return heading_cn(doc, text, 12.5)

def eq(img, num):
    builder = EQMAP.get(img.replace(".png", ""))
    if builder is None:
        raise ValueError("未知公式: " + img)
    return add_equation(doc, builder(), num)

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# 路径自解析：仓库克隆/解压到任意目录均可运行
BASE = os.path.dirname(os.path.abspath(__file__))
# 作者身份信息取自本地私有模块 author_local.py（已 .gitignore 排除）；缺失时用中性占位符
import sys as _sys
_sys.path.insert(0, BASE)
try:
    from author_local import AUTHOR_CN as _AUTHOR_CN, AFFIL_CN as _AFFIL_CN, EMAIL as _EMAIL
except ImportError:
    _AUTHOR_CN, _AFFIL_CN, _EMAIL = "[作者]", "[单位]", "[email]"
FIG = os.path.join(BASE, "figures")
OUT = os.path.join(BASE, "out", "论文_CN_latest.docx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

def set_cn(run):
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

def para(text, size=12, bold=False, align=None, indent=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold
    set_cn(r)
    if align is not None: p.alignment = align
    if indent and align is None: p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    return p

def fig(img, caption, w=14.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(FIG, img), width=Cm(w))
    cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(caption); r.font.size = Pt(9.5); r.bold = True; set_cn(r)


def _vcenter_content(table):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def table(headers, rows, caption):
    para(caption, 10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]; c.paragraphs[0].add_run(htxt).bold = True
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].runs[0].font.size = Pt(9)
        set_cn(c.paragraphs[0].runs[0])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(9); set_cn(r)
    make_three_line(t)
    _vcenter_content(t)
    return t

# ================= 标题 =================
para("祁连山区域生态安全预警：时序动态、预测与空间驱动因子", 17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para(_AUTHOR_CN, 11, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para(_AFFIL_CN, 10, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para("* 通讯作者：" + _EMAIL, 10, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
doc.add_paragraph()

# ================= 摘要 =================
h1("摘要")
para("自然保护区对生物多样性保护至关重要，但其管理往往缺乏能将监测数据转化为可操作信号的定量预警工具。本研究结合驱动力-压力-状态-影响-响应（DPSIR）指标体系、主客观组合赋权、灰色预测与地理探测器归因，仅使用公开气候（CRU TS）与土地覆被（CLCD）数据，对中国祁连山区域（含祁连山国家公园）2005-2023 年的生态安全（ES）进行评价与预测。区域生态安全指数（ESI）于 2010 年升至峰值 0.709（较安全），此后回落并在轻警区间趋于稳定，2023 年为 0.486；19 年中有 16 年为轻警等级，且 2023 年等级对主客观赋权比例不敏感。经三模型比较与滚动回测选定的非线性灰色伯努利模型给出的 2024-2028 年预测显示，生态安全将在轻警等级内继续缓慢下降。空间上，裸地/退化占比（q=0.80）与草地占比（q=0.71）主导生态安全的空间异质性（均 p<0.001）：决定格局的是干旱程度与植被覆盖，而非气温。据此提出分区管理建议——加强轻警区草地恢复、规范耕地扩张——并提供可推广至其他保护区的完全可复现流程。")
para("关键词：生态安全；预警；DPSIR；GM(1,1)；地理探测器；祁连山区域", 10, bold=True, indent=False)
doc.add_paragraph()

# ================= 1 引言 =================
h1("1  引言")
h2("1.1  背景")
para("全球生物多样性正以前所未有的速度下降：自 1970 年以来脊椎动物种群规模平均下降 68%，晚更新世以来陆生哺乳动物生物量至少减少 85%，非法野生动物贸易年价值高达 230 亿美元 (WWF, 2020; IPCC, 2014)。栖息地丧失、气候变化与偷猎共同威胁着生态系统稳定。保护区已成为全球就地生物多样性保护的基石。")
para("然而，保护区的有效性受制于有限的监测能力。常规野外调查劳动密集、成本高昂，难以在广大或难以到达的区域持续开展；红外相机与遥感提供了更丰富的数据，但需要系统的分析框架将原始观测转化为管理决策 (Kerry et al., 2022; Bijl and Heltai, 2022)。一个反复出现的缺口是缺乏“预警”信息：生态退化往往在已变得严重之后才被发现，留给预防性干预的空间很小。")
para("这一“监测-预警”缺口已被多项保护倡议反复指出。例如，国际社会呼吁通过建立融合传感器网络、遥感与人工智能的云监测平台来减少非法野生动物贸易，实时发出栖息地威胁警报 (WWF, 2020; Kerry et al., 2022)；此类平台的预警决策模块通常依托于复合生态安全指标，将多维环境信息压缩为可操作信号。用真实数据对预警模块进行形式化与实证验证——而不是停留在平台蓝图层面——因此是对保护监测议程的直接贡献。")
para("生态安全（ES）视角提供了这样一种系统化的形式化路径。生态安全评价将多维环境、社会与经济信息压缩为可在时空上追踪的复合指数，为分区、监测与预警提供定量基础 (Liu et al., 2023, 2022)。本研究采用该视角，并将其应用于栖息地完整性直接决定野生动物保护成效的国家公园。")
h2("1.2  研究现状与知识缺口")
para("PSR 与 DPSIR 框架被广泛用于组织生态安全指标体系 (Hou et al., 2024; Yang et al., 2023)。主观赋权（AHP）(Saaty, 1980)、客观赋权（熵权法）及其组合是主流方法。灰色模型（GM(1,1)）(Deng, 1982) 适用于小样本生态安全时间序列；Liu 等将 DPSIR、组合赋权、GM(1,1) 与地理探测器集成到保护区尺度的预警框架中 (Liu et al., 2023)。地理探测器方法 (Wang et al., 2016; Wang and Hu, 2012) 可量化各因子对空间生态安全异质性的贡献 (Lyu et al., 2026; Zhang et al., 2024; Gao et al., 2025)。")
para("现有研究仍存在知识缺口：（1）保护区尺度研究通常孤立地应用 DPSIR–GM(1,1)–GeoDetector 各组件，或采用单一赋权方案；（2）生态安全结果对权重比例 α 的敏感性很少被检验；（3）GM(1,1) 对非单调生态安全序列的适用性很少被诊断；（4）很少有研究将预警结果转化为具体的分区管理建议。")
h2("1.3  本文贡献")
para("（1）构建了面向国家公园区域的开放数据预警管线，将基于 DPSIR 的综合指数与预测、空间归因相耦合，并包含单调性诊断与三模型预测对比；（2）进行了权重比例敏感性分析（α=0-1），证明预警等级对主客观赋权平衡具有稳健性；（3）为祁连山区域（含祁连山国家公园）2005-2023 年提供了真实数据证据链（时间趋势、稳健性、预测、空间驱动因子）及 2024-2028 五年预测；（4）提出了分区管理建议，并提供完全可复现的数据与代码。")

# ================= 2 研究区与数据 =================
h1("2  研究区与数据")
h2("2.1  研究区概况")
para("祁连山区域以祁连山国家公园为核心，位于青藏高原东北缘，横跨甘肃与青海两省交界（图 1）。公园保护着典型的山地生态系统——冰川、雪原、高山草甸与荒漠草原——为河西走廊及内陆河流域提供关键水源服务。由于权威的公园边界矢量未公开获取，本文分析范围取覆盖公园及其缓冲县的矩形区域（36.0-40.5°N，93.5-103.5°E），因此结论在区域尺度上解读。研究区 2005-2023 年年均气温 0.8-2.0°C、年降水量 167-284 mm。分析单元包括整个区域及 11 个县（天祝、肃南、肃北、阿克塞、民乐、山丹、门源、祁连、刚察、德令哈、天峻）。")
fig("fig1_study_area.png", "图1  祁连山区域地形、祁连山国家公园及县域位置", w=11.5)
h2("2.2  数据来源")
table(["数据类型", "指标", "来源", "时段"],
      [["气候", "降水量（mm）、年均气温（°C）", "CRU TS 4.10 (Harris et al., 2020)", "2005-2023"],
       ["土地覆被", "草地、森林、水域、耕地、裸地占比", "CLCD 30 m (Yang and Huang, 2021)", "2005-2023"]],
      "表1  本研究使用的数据来源")
para("本研究使用的数据源汇总于表 1。土地覆被按 2005、2010、2015、2020、2023 五个时期计算并线性插值到年值；气候格网聚合到区域尺度与 11 个县域质心，因此县域数值为质心窗口估计而非面积加权统计。")

# ================= 3 方法 =================
h1("3  生态安全预警模型")
h2("3.1  DPSIR 框架与指标体系")
para("遵循 DPSIR 框架 (Hou et al., 2024; Yang et al., 2023)，九项指标覆盖各准则层：驱动力（降水、气温）、压力（裸地/退化占比、耕地占比）、状态（草地、森林、水域占比）、影响（草地变化率）、响应（生态用地占比）（表 2）。气温与降水取正向：在寒冷干旱的高寒环境中，热量与水分为植被生长的共同限制因子，升温带来的蒸散发增强等相反效应在空间归因分析中讨论。分析框架见图 2。")
fig("fig_dpsir_framework.png", "图2  分析框架：DPSIR 因果链与 ESI 预警管线（虚线箭头为管理反馈）", w=16.0)
table(["准则层", "指标", "单位", "方向"],
      [["驱动力 D", "年降水量", "mm", "+"],
       ["驱动力 D", "年均气温", "°C", "+"],
       ["压力 P", "裸地/退化占比", "%", "−"],
       ["压力 P", "耕地占比", "%", "−"],
       ["状态 S", "草地占比", "%", "+"],
       ["状态 S", "森林占比", "%", "+"],
       ["状态 S", "水域占比", "%", "+"],
       ["影响 I", "草地变化率", "%/年", "+"],
       ["响应 R", "生态用地占比", "%", "+"]],
      "表2  DPSIR 指标体系（+/− 表示对生态安全的正/负向）")
h2("3.2  组合赋权（AHP + 熵权）")
para("熵权：ej = −(1/ln n) Σi pij ln pij，pij = x′ij / Σi x′ij，wEj = (1 − ej) / Σj(1 − ej)。AHP 权重采用主特征向量法：准则层判断矩阵（5 个准则，Saaty 1-9 标度；表 3）依 DPSIR 层次优先级（状态 > 响应 > 影响 > 驱动力 = 压力）构造，得 λmax = 5.015、CI = (λmax − n)/(n − 1) = 0.0038、CR = CI/RI = 0.0034 < 0.1（n = 5 时 RI = 1.12），满足一致性检验 (Saaty, 1980)；各准则层内指标均分该层权重。组合权重为：")
eq("eq_weight", 1)
para("以 α = 0.4 为基线，并在 α = 0-1 全区间进行敏感性分析。")
table(["准则层", "D", "P", "S", "I", "R", "权重"],
      [["驱动力 D", "1", "1", "1/5", "1/2", "1/3", "0.081"],
       ["压力 P", "1", "1", "1/5", "1/2", "1/3", "0.081"],
       ["状态 S", "5", "5", "1", "3", "2", "0.438"],
       ["影响 I", "2", "2", "1/3", "1", "1/2", "0.149"],
       ["响应 R", "3", "3", "1/2", "2", "1", "0.250"]],
      "表3  AHP 准则层判断矩阵（Saaty 1-9 标度）与准则层权重")
h2("3.3  生态安全指数（ESI）")
para("指标按方向标准化（式 2），正向指标越高生态越安全，负向指标反之：", 11)
eq("eq_norm", 2)
eq("eq_esi", 3)
h2("3.4  预警预测模型与等级划分")
para("GM(1,1) 模型 (Deng, 1982) 基于一次累加序列，以最小二乘估计参数（â, b̂）：")
eq("eq_gm", 4)
para("GM(1,1) 假设单调指数趋势，而观测 ESI 序列在 2010 年见顶后回落、违背了这一假设，故本文诊断该非单调性并额外采用广义灰色伯努利模型（NGBM，幂指数 γ）与 ARIMA 对比。模型精度由后验差比值 C = S2/S1 与小误差概率 P 评价：优（C<0.35，P>0.95）、合格（C<0.50，P>0.80）、勉强合格（C<0.65，P>0.70）。预警等级：安全 [0.8,1]、较安全 [0.6,0.8)、轻警 [0.4,0.6)、中警 [0.2,0.4)、重警 [0,0.2)；等级界限在归一化区间 [0,1] 上按 0.2 等距划分，与 DPSIR 类保护区生态安全预警框架的五级划分惯例一致 (Liu et al., 2023)，等距划分便于跨研究比较并避免样本相关的断点，后文将表明观测序列的预警等级对权重比例不敏感。模型精度另由滚动回测验证：以 2005-2018 年拟合、以留出的 2019-2023 年观测评估。")
h2("3.5  空间驱动因子识别（GeoDetector）")

eq("eq_q", 5)

# ================= 4 结果 =================
h1("4  结果与分析")
fig("fig_indicator_trends.png", "图3  九项 DPSIR 指标变化趋势（2005-2023；圆点为五个土地覆被观测时点，虚线为 2010 年即 ESI 峰值年）", w=15.2)
para("图 3 展示 ESI 背后的指标原始动态：降水量在 166.9-284.4 mm/年之间波动，由 2005 年的 228.0 mm 降至 2023 年的 166.9 mm；年均气温由 1.29°C 升至 1.92°C。土地覆被类指标（由五个 CLCD 时点线性插值）变化缓慢：生态用地占比（草地+森林+水域）由 2005 年的 37.9% 升至 2010 年的 39.6%，2023 年为 39.0%；裸地/退化占比先降（58.2%→56.5%）后微升（57.0%）；水域（1.75%-2.12%）与森林（1.15%-1.29%）占比稳步上升，耕地小幅增加（3.8%-4.0%）；草地变化率波动最大（-0.47% 至 +0.84%/年）。这些序列经标准化与组合赋权后即构成下文分析的 ESI 轨迹（图 4）。")
h2("4.1  指标权重")
table(["指标", "方向", "熵权", "AHP", "组合"],
      [["降水（mm）", "+", "0.111", "0.041", "0.083"],
       ["气温（°C）", "+", "0.063", "0.041", "0.054"],
       ["草地占比", "+", "0.075", "0.146", "0.103"],
       ["森林占比", "+", "0.103", "0.146", "0.120"],
       ["水域占比", "+", "0.115", "0.146", "0.128"],
       ["裸地/退化占比", "−", "0.058", "0.041", "0.051"],
       ["耕地占比", "−", "0.171", "0.041", "0.119"],
       ["草地变化率", "+", "0.245", "0.149", "0.206"],
       ["生态用地占比", "+", "0.059", "0.250", "0.136"]],
      "表4  指标权重（熵权、AHP 与组合权重，α=0.4）")
para("熵权（表 4）识别出草地变化率（w=0.245）与耕地占比（w=0.171）为信息量最大的信号；组合权重（α=0.4）下草地变化率（0.206）、生态用地占比（0.136）与水域占比（0.128）权重最高，表明草地动态、生态用地规模与水资源共同主导生态安全的时序变化（图 7）。")
h2("4.2  ESI 时序与预警等级")
fig("fig3_esi_timeseries.png", "图4  观测 ESI（2005-2023）与预测（2024-2028）及预警等级带", w=14.4)
para("区域 ESI 从 2005 年的 0.258（中警）升至 2010 年的峰值 0.709（较安全），随后回落并在 0.45-0.59（轻警）间波动，2023 年为 0.486。Mann-Kendall 趋势检验表明，2005-2023 年 ESI 的下降趋势不显著（Z=-1.120，p=0.263），Sen 斜率为 -0.0032/年，说明区域处于稳定但脆弱的轻警状态，而非已确认的持续退化。")
h2("4.3  预测对比：GM(1,1)、NGBM 与 ARIMA")
fc = pd.read_csv(os.path.join(BASE, "data", "results_forecast_comparison.csv"))
table(["模型", "RMSE", "MAE", "C", "P", "2024 预测", "2028 预测"],
      [["GM(1,1)", 0.052, 0.035, 0.597, 0.789, f'{fc["gm11"].iloc[0]:.3f}', f'{fc["gm11"].iloc[-1]:.3f}'],
       ["NGBM (γ=0.20)", 0.038, 0.029, 0.442, 0.895, f'{fc["ngbm"].iloc[0]:.3f}', f'{fc["ngbm"].iloc[-1]:.3f}'],
       ["ARIMA(0,1,1)", 0.098, 0.067, 1.086, 0.632, f'{fc["arima"].iloc[0]:.3f}', f'{fc["arima"].iloc[-1]:.3f}']],
      "表5  2005-2023 ESI 序列的预测模型对比")
para("注：精度等级按灰色模型惯例判定——合格（C<0.50，P>0.80）、勉强合格（C<0.65，P>0.70）、其余为不合格；本文采用 NGBM。", 9, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para("2024-2028 年 ESI 预测在无干预情景下由 0.446 缓慢降至 0.394，参数化 95% 预测区间分别为 2024 年 0.358-0.533、2028 年 0.306-0.482，跨越中警至轻警等级，表明预测不确定性不可忽略（表 5；图 4）。该预测由非线性灰色伯努利模型（NGBM）给出：在三个候选模型中（表 5），NGBM 达到合格精度（C=0.44，P=0.89），优于 GM(1,1)（C=0.60，P=0.79，勉强合格）与 ARIMA（C=1.09，不合格）；以 2005-2018 年拟合、2019-2023 年留出检验的滚动原点回测进一步确认其优势（RMSE 0.034，GM(1,1) 为 0.065，改善 47%），支持非线性灰色模型用于小样本、非平稳的保护区尺度生态安全序列 (Althobaiti and Shabri, 2022)。")
h2("4.4  权重比例 α 的稳健性")
alpha_df = pd.read_csv(os.path.join(BASE, "data", "results_alpha_sensitivity.csv"))
fig("fig4_alpha_sensitivity.png", "图5  2023 年 ESI 对权重比例 α 的敏感性", w=11.5)
para(f"α=0-1 全区间敏感性分析表明，2023 年 ESI 由 α=0（纯熵权）的 {alpha_df['esi_2023'].min():.3f} 升至 α=1（纯 AHP）的 {alpha_df['esi_2023'].max():.3f}，但预警等级在所有 α 取值下均为「轻警」（图 5）。19 年序列中轻警年数由 α=0.3-0.4 时的 16 年降至 α≥0.9 时的 5 年，且 α≥0.7 时出现 1 个重警年；基线 α=0.4 处于最稳定区间。当前状态的预警结论对主客观赋权平衡具有稳健性，但 AHP 主导的权重会放大历史等级的分散度。")
h2("4.5  空间分布、驱动因子与障碍因子")
table(["县域", "ESI（2020）", "县域", "ESI（2020）"],
      [["肃南", "0.466", "刚察", "0.289"],
       ["天祝", "0.395", "山丹", "0.260"],
       ["祁连", "0.390", "天峻", "0.251"],
       ["德令哈", "0.366", "民乐", "0.166"],
       ["肃北", "0.344", "阿克塞", "0.074"],
       ["门源", "0.306", "", ""]],
      "表6  县域静态 ESI（2020）")
fig("fig5_spatial_esi.png", "图6  2020 年点尺度 ESI 空间分布", w=12.8)
para("县域静态 ESI（2020，采用与区域分析相同的组合权重）以东部的肃南（0.466）与天祝（0.395）最高、西部的阿克塞（0.074）最低（表 6；图 6）。草地覆盖较高的东部县域（肃南、天祝、祁连）高于干旱的西部县域（阿克塞、德令哈），但格局受局地土地覆被构成调节（如以耕地为主的民乐低于部分西部县域）。该梯度表明保护资源与预警优先级应在区域内部差异化配置。")
para("对 1,914 个格网样本的地理探测器分析识别出裸地/退化占比（q=0.80，p<0.001）为主导空间驱动因子，其次是草地占比（q=0.71，p<0.001）与降水（q=0.51，p<0.001）；气温（q=0.17，p<0.001）与森林占比较弱（置换检验 999 次；图 8）。干旱化与植被覆盖——而非气温——控制着祁连山区域生态安全的空间格局。交互探测器进一步表明所有因子对均呈“增强”交互：气温×裸地（q=0.849）、降水×裸地（q=0.833）与草地×裸地（q=0.818）的联合效应均超过各自单独贡献（p<0.001），说明热-旱与水分-植被耦合叠加了空间生态安全异质性（表 7；图 9）。")
fig("fig_weights.png", "图7  熵权、AHP 与组合权重（α=0.4）对比", w=13.5)
fig("fig_geodetector_q.png", "图8  单因子探测 q 值（2020；置换检验 999 次，均 p<0.001）", w=10.5)
fig("fig_interaction.png", "图9  六对因子的交互探测 q 值（q1/q2 单因子，q_int 联合；均呈增强）", w=14.2)
table(["因子对", "q1", "q2", "q_int", "交互关系"],
      [["草地 × 裸地", "0.708", "0.802", "0.818", "增强"],
       ["降水 × 裸地", "0.506", "0.802", "0.833", "增强"],
       ["降水 × 草地", "0.506", "0.708", "0.808", "增强"],
       ["气温 × 裸地", "0.170", "0.802", "0.849", "增强"],
       ["气温 × 草地", "0.170", "0.708", "0.723", "增强"],
       ["气温 × 降水", "0.170", "0.506", "0.529", "增强"]],
      "表7  地理探测器交互探测（2020）")
para("2020 年点尺度 ESI 的 Moran's I 为 0.643（p=0.005，置换检验；1,914 个样本，反距离 k 近邻权重），表明空间格局显著聚集，支持地理探测器的归因。")
para("障碍因子分析（沿用复合指数生态安全评价中的障碍度诊断 (Jing et al., 2024)）显示约束条件存在阶段性转变：恢复期（2005-2010）主要障碍为水域占比（0.233）、森林占比（0.221）与生态用地占比（0.151），反映水资源与森林规模不足；平台期（2011-2023）转为草地变化率（0.364）与耕地占比（0.206），说明草地动态与耕地扩张已成为主导约束，与权重结果一致（图 10）。")
fig("fig_obstacle.png", "图10  九项指标逐年障碍度构成（2005-2023，每年各行之和为 1）", w=14.6)

# ================= 5 讨论 =================
h1("5  讨论")
h2("5.1  生态安全变化及其驱动因素")
para("2005-2010 年 ESI 回升与大规模生态修复（退牧还草、植被恢复）及丰水年份一致；此后的轻警平台反映了裸地主导的土地覆被（约 56%）与气候波动的持续影响。草地动态的主导权重印证了高寒草甸生态系统对区域生态安全的重要性，与高山地区草地退化是生态安全首要约束的既有结论一致 (Lyu et al., 2026; Gao et al., 2025)。东西梯度与地理探测器结果表明，干旱化与植被覆盖——而非气温——控制着空间异质性 (Zhang et al., 2024)。")
para("对野生动物与生物多样性保护而言，轻警状态与预测下降传递了具体信息：祁连山区域栖息地完整性并未改善，依赖静态栖息地评估的监测项目可能漏判缓慢退化。本文形式化的 ESI 预警信号正是保护监测平台所需的决策支持输入——将原始环境数据与可触发管理行动的预警等级相连接 (Kerry et al., 2022; Bijl and Heltai, 2022)。该框架源于国际野生动物保护监测倡议的预警决策模块：通过复合生态安全指标，将遥感与地面监测数据转化为可操作的预警等级，支撑栖息地监测-预警-响应闭环；本文以祁连山国家公园为例完成了从框架到实证的验证。")
h2("5.2  与既有研究对比")
para("相对 Liu et al. (2023) 提出的保护区尺度预警框架（其组合同样包含 DPSIR、组合赋权、GM(1,1) 与地理探测器），本文结果确认该工具链可迁移至国家公园区域，并在表 8 所列三方面有所扩展。其一，预警等级在权重比例 α 的全区间内保持不变，而同类 DPSIR 评价多采用单一赋权方案 (Hou et al., 2024; Yang et al., 2023)——当前状态的预警结论因此不依赖于分析者的赋权选择。其二，实测 ESI 序列非单调，违背 GM(1,1) 的单调趋势假设；非线性灰色伯努利模型以合格精度刻画了这一行为，而 ARIMA 不合格，这与非线性灰色变式适用于短小非平稳序列的证据一致 (Althobaiti and Shabri, 2022)。")
para("其三，将区域时间序列与格网、县域尺度归因相衔接，揭示出单一尺度研究无法显现的尺度依赖性：主导时间变化的指标（草地变化率与耕地占比；表 4）与主导空间异质性的指标（裸地/退化占比与草地占比；图 8）并不相同。预警设计因此应将草地与耕地动态的时间监测、土地退化的空间精准管控配对，而非假设同一组驱动因子同时适用于两个尺度。")
table(["方法特征", "Liu 等 (2023)", "长白山 DPSIRM", "DPSIR-DEA", "本研究"],
      [["DPSIR/PSR 框架", "是", "是", "是", "是"],
       ["组合赋权", "是", "仅 AHP", "熵权-DEA", "是，含 α 敏感性"],
       ["GM(1,1) 预测", "是", "否", "否", "是，对比 NGBM/ARIMA"],
       ["地理探测器归因", "是", "否", "否", "是，含交互探测"],
       ["县域空间分析", "否", "是", "是", "是"],
       ["权重稳健性检验", "否", "否", "否", "是"],
       ["开放数据与代码", "否", "否", "否", "是"]],
      "表8  本研究框架与代表性保护区尺度生态安全研究的方法特征对比")
h2("5.3  管理启示")
para("（1）轻警区应优先实施草地恢复与放牧调控；干旱的西部县（阿克塞、德令哈）需要荒漠化监测与水资源管理而非单纯的植被恢复。（2）农牧交错带的耕地扩张应加以控制。（3）五级预警分类可支撑与祁连山国家公园功能分区衔接的 ESI 阈值触发响应机制，使预警等级跨越边界时自动提升监测频率与管理关注——将预警概念落实为栖息地与野生动物保护的可操作机制。")
h2("5.4  局限与展望")
para("局限包括：区域聚合（县域 ESI 为基于质心窗口的静态评价，待更精细行政边界）、土地覆被在五个时期之间插值、灰色模型对非单调序列精度中等（已通过 NGBM/ARIMA 对比部分缓解），且 ESI 尚未与独立外部生态指标（如 NDVI、草地覆盖地面观测）交叉验证。未来工作应（i）纳入更高时间分辨率的土地覆被并实现近实时预警更新，（ii）将 ESI 与野生动物种群及栖息地质量指标耦合，（iii）将框架扩展至多保护区对比以检验可迁移性，（iv）以独立外部指标对 ESI 进行交叉验证。")

# ================= 6 结论 =================
h1("6  结论")
para("仅使用公开气候与土地覆被数据，本文对祁连山区域 2005-2023 年的生态安全进行评价与预测，管线完全可复现。2005-2023 年区域 ESI 介于 0.26-0.71；2011 年以来区域持续处于轻警状态（2023 年 ESI=0.486），Mann-Kendall 检验表明近年下降趋势不显著（p=0.263），生态安全呈稳定但脆弱的态势。")
para("2024-2028 年 NGBM 预测由 0.446 缓慢降至 0.394（不干预场景），需主动干预；参数化 95% 区间为 0.358-0.533（2024）与 0.306-0.482（2028），C=0.44 达合格精度，滚动回测 RMSE=0.034，预警等级存在不确定性。预警等级对主客观赋权比例稳健，草地动态与耕地压力为主导信号；地理探测器识别出干旱化（裸地占比）为空间主导驱动因子。建议将 ESI 预警纳入祁连山国家公园适应性管理，优先针对障碍因子（草地动态、耕地扩张）实施草地恢复与耕地管控；该框架可推广至其他自然保护区，为野生动物栖息地保护监测提供预警-响应决策支持。")

# ================= 数据、声明与参考文献 =================
doc.add_paragraph()
h1("作者贡献")
para("作者构思并设计了本研究，完成数据分析与建模、结果解释与论文撰写。")
h1("数据可获得性")
para("气候数据公开获取自 CRU TS 4.10（https://crudata.uea.ac.uk/cru/data/hrg/）；土地覆被数据来自 CLCD 数据集（https://doi.org/10.5194/essd-13-3907-2021）；高程数据来自 Copernicus DEM。支撑本研究的处理后数据集与分析代码公开于 https://github.com/mllbmy/qilian-esi-early-warning。")
h1("利益竞争")
para("利益冲突声明：作者声明不存在任何利益冲突。")
h1("伦理声明")
para("本研究不涉及人类受试者与动物实验。所用数据集均为公开获取，研究遵循期刊的伦理标准。")
h1("基金资助")
para("本研究未获得任何公共、商业或非营利机构资助项目的专项经费支持。")
h1("生成式人工智能与人工智能辅助技术使用声明")
para("在本工作准备过程中，作者使用 DeepSeek（人工智能语言模型）辅助稿件的撰写、结构组织与语言润色。使用该工具后，作者对内容进行了必要的审阅与修改，并对本文发表内容的真实性、准确性与完整性承担全部责任。")
doc.add_paragraph()
h1("参考文献")
refs = [
        "Althobaiti, S., Shabri, A., 2022. Prediction of CO2 emissions in Saudi Arabia using nonlinear grey Bernoulli model NGBM(1,1) compared with GM(1,1) model, in: Journal of Physics: Conference Series, IOP Publishing. p. 012011. doi: 10.1088/1742-6596/2259/1/012011.",
        "Bijl, H., Heltai, M., 2022. A narrative review on the use of camera traps and machine learning in wildlife research. Columella – Journal of Agricultural and Environmental Sciences 9, 47–69. doi: 10.18380/szie.colum.2022.9.2.47.",
        "Deng, J., 1982. Control problems of grey systems. Systems & Control Letters 1, 288–294.",
        "Gao, J., Pan, N., Zhou, D., 2025. Integrating landscape ecological risk and ecosystem services for ecological zoning in the Qilian Mountain National Park. Ecosystem Health and Sustainability 11, 0441. doi: 10.34133/ehs.0441.",
        "Harris, I., Osborn, T.J., Jones, P., Lister, D., 2020. Version 4 of the CRU TS monthly high-resolution gridded multivariate climate dataset. Scientific Data 7, 109. doi: 10.1038/s41597-020-0453-3.",
        "Hou, M., Li, L., Yu, H., Jin, R., Zhu, W., 2024. Ecological security evaluation of wetlands in Changbai Mountain area based on DPSIRM model. Ecological Indicators 160, 111773. doi: 10.1016/j.ecolind.2024.111773.",
        "IPCC, 2014. Climate Change 2014: Impacts, Adaptation, and Vulnerability. Part A: Global and Sectoral Aspects. Cambridge University Press.",
        "Jing, X., Tao, S., Hu, H., Sun, M., Wang, M., 2024. Spatio-temporal evaluation of ecological security of cultivated land in China based on DPSIR-entropy weight TOPSIS model and analysis of obstacle factors. Ecological Indicators 166, 112579. doi: 10.1016/j.ecolind.2024.112579.",
        "Kerry, R.G., Montalbo, F.J.P., Das, R., Patra, S., Mahapatra, G.P., Maurya, G.K., Nayak, V., Jena, A.B., Ukhurebor, K.E., Jena, R.C., Gouda, S., Majhi, S., Rout, J.R., 2022. An overview of remote monitoring methods in biodiversity conservation. Environmental Science and Pollution Research 29, 80179–80221. doi: 10.1007/s11356-022-23242-y.",
        "Liu, J., Cao, X., Zhao, L., Dong, G., Jia, K., 2022. Spatiotemporal differentiation of land ecological security and its influencing factors: A case study in Jinan, Shandong Province, China. Frontiers in Environmental Science 10, 824254. doi: 10.3389/fenvs.2022.824254.",
        "Liu, Y., Wang, C., Wang, H., Chang, Y., Yang, X., Zang, F., 2023. An integrated ecological security early-warning framework in the national nature reserve based on the gray model. Journal for Nature Conservation 73, 126394. doi: 10.1016/j.jnc.2023.126394.",
        "Lyu, X., Li, X., Wang, K., Cao, W., Zhang, C., 2026. Anthropogenic activities amplify spatiotemporal variations in regional ecological security patterns dominated by natural factors: Evidence from the West Liaohe River Basin, China. Journal of Arid Land 18, 752–773. doi: 10.1016/j.jaridl.2026.05.002.",
        "Saaty, T.L., 1980. The Analytic Hierarchy Process. McGraw-Hill, New York.",
        "Wang, J.F., Hu, Y., 2012. Environmental health risk detection with geogdetector. Environmental Modelling & Software 33, 114–115. doi: 10.1016/j.envsoft.2012.01.015.",
        "Wang, J.F., Zhang, T.L., Fu, B.J., 2016. A measure of spatial stratified heterogeneity. Ecological Indicators 67, 250–256. doi: 10.1016/j.ecolind.2016.02.052.",
        "WWF, 2020. Living Planet Report 2020 – Bending the curve of biodiversity loss. Technical Report. WWF International. Gland, Switzerland.",
        "Yang, G., Gui, Q., Liu, J., Chen, X., Cheng, S., 2023. Spatial-temporal evolution and driving factors of ecological security in China based on DPSIR-DEA model: A case study of the Three Gorges reservoir area. Ecological Indicators 154, 110777. doi: 10.1016/j.ecolind.2023.110777.",
        "Yang, J., Huang, X., 2021. The 30 m annual land cover dataset and its dynamics in China from 1990 to 2019. Earth System Science Data 13, 3907–3925. doi: 10.5194/essd-13-3907-2021.",
        "Zhang, W., Liu, Z., Qin, K., Dai, S., Lu, H., Lu, M., 2024. Long-term dynamic monitoring and driving force analysis of eco-environmental quality in China. Remote Sensing 16, 1028. doi: 10.3390/rs16061028.",
        ]
for r in refs:
    para(r, 9.5, indent=False)

doc.save(OUT)
print("OK:", OUT)
