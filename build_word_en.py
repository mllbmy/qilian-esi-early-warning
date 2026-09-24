# -*- coding: utf-8 -*-
"""从 manuscript.tex 生成英文对照 Word（与 CN Word 同一套版式；投稿 PDF 不动）

结构镜像 CN Word：标题块 → Abstract(+Keywords) → §1..§6 → Highlights →
声明段（Data availability / CRediT / Competing interest / Funding）→ References
图/表/公式编号与 CN Word 及投稿 PDF 完全一致（6 图 / 8 表 / 5 式）。
"""
import io
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from word_style import (add_equation, make_three_line, heading_cn,
                        EQNORM_EN, EQ3, EQ4, EQ5, EQ6)

# 路径自解析：仓库克隆/解压到任意目录均可运行
BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(BASE, "figures")
OUT = os.path.join(BASE, "out", "论文_EN_review.docx")
LEFT = []


def rd(p):
    return io.open(p, encoding="utf-8", errors="replace").read() if os.path.isfile(p) else ""


TEX = rd(os.path.join(BASE, "manuscript.tex"))
AUX = rd(os.path.join(BASE, "manuscript.aux"))
BBL = rd(os.path.join(BASE, "manuscript.bbl"))

BIB = {m.group(1): (m.group(4), m.group(3)) for m in re.finditer(
    r"\\bibcite\{([^}]+)\}\{\{(\d+)\}\{(\d{4})\}\{\{([^}]*)\}\}", AUX)}
# 声明段：由尾段统一输出，正文遍历遇到时必须跳过（否则 Word 里出现两组）
BACK_SECTIONS = ("Data availability", "CRediT authorship contribution statement",
                 "Declaration of competing interest", "Funding",
                 "Declaration of generative AI and AI-assisted technologies in the writing process")

REFS = {}
for m in re.finditer(r"\\newlabel\{([^}]+)\}\{\{([^}]+)\}", AUX):
    REFS.setdefault(m.group(1), m.group(2))
BBL_KEYS = re.findall(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}", BBL)

GREEK = {r"\alpha": "α", r"\beta": "β", r"\gamma": "γ", r"\delta": "δ", r"\Delta": "Δ",
         r"\sigma": "σ", r"\mu": "μ", r"\lambda": "λ", r"\theta": "θ", r"\pi": "π",
         r"\rho": "ρ", r"\tau": "τ", r"\phi": "φ", r"\omega": "ω", r"\Omega": "Ω",
         r"\times": "×", r"\pm": "±", r"\approx": "≈", r"\leq": "≤", r"\geq": "≥",
         r"\le": "≤", r"\ge": "≥", r"\sim": "∼", r"\circ": "°", r"\rightarrow": "→",
         r"\to": "→", r"\cdot": "·", r"\infty": "∞", r"\in": "∈", r"\sum": "Σ",
         r"\S": "§", r"\%": "%", r"\&": "&", r"\_": "_", r"\,": " ", r"\;": " ", r"\!": "",
         r"\ ": " ", r"\quad": " ", r"\qquad": " "}


_SUPMAP = str.maketrans("0123456789+-=()ni", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿⁱ")
_SUPSET = set("⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿⁱ")


def _sup(x):
    """斜率单位之类的内联上标：可映射则用 Unicode 上标，否则退回纯文本"""
    y = x.translate(_SUPMAP)
    return y if y and all(c in _SUPSET for c in y) else x


_HATMAP = {"a": "â", "b": "b̂", "e": "ê", "i": "î", "n": "n̂", "o": "ô", "u": "û", "x": "x̂", "y": "ŷ", "w": "ŵ", "z": "ẑ"}


def math2txt(m):
    m = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"\1/\2", m)
    m = re.sub(r"\\dfrac\{([^{}]*)\}\{([^{}]*)\}", r"\1/\2", m)
    for _ in range(2):
        m = re.sub(r"\\text\{([^{}]*)\}", r"\1", m)
        m = re.sub(r"\\mathrm\{([^{}]*)\}", r"\1", m)
        m = re.sub(r"\\emph\{([^{}]*)\}", r"\1", m)
    m = m.replace(r"^{\circ}", "°")
    m = re.sub(r"\^\{([^{}]*)\}", lambda x: _sup(x.group(1)), m)
    m = re.sub(r"_\{([^{}]*)\}", r"\1", m)
    m = re.sub(r"\^(\w)", lambda x: _sup(x.group(1)), m)
    m = re.sub(r"_(\w)", r"\1", m)
    m = re.sub(r"\\hat\s*\{?([a-zA-Z])\}?", lambda x: _HATMAP.get(x.group(1), x.group(1)), m)
    for k, v in GREEK.items():
        m = m.replace(k, v)
    m = m.replace("{", "").replace("}", "").replace("\\", "")
    return m


def cite(keys):
    """\\citep{a,b} -> (Author, Year; Author, Year)，同作者合并年份（elsarticle-harv 体例）"""
    order, by = [], {}
    for k in keys.split(","):
        k = k.strip()
        if not k:
            continue
        a, y = BIB.get(k, ("?", "?"))
        a = a.replace("~", " ")
        if a not in by:
            by[a] = []
            order.append(a)
        by[a].append(y)
    if not order:
        return ""
    return "(" + "; ".join("%s, %s" % (a, ", ".join(by[a])) for a in order) + ")"


def tex2txt(s):
    s = re.sub(r"\\label\{[^}]*\}", "", s)
    s = re.sub(r"\\citep\{([^}]*)\}", lambda m: cite(m.group(1)), s)
    s = re.sub(r"\\citet\{([^}]*)\}", lambda m: cite(m.group(1)), s)
    s = re.sub(r"\\cite\{([^}]*)\}", lambda m: cite(m.group(1)), s)
    s = re.sub(r"\\eqref\{([^}]*)\}", lambda m: "(" + REFS.get(m.group(1), "?") + ")", s)
    s = re.sub(r"\\ref\{([^}]*)\}", lambda m: REFS.get(m.group(1), "?"), s)
    s = re.sub(r"\$([^$]*)\$", lambda m: math2txt(m.group(1)), s)
    s = re.sub(r"\\multirow\{\d+\}\{[^}]*\}\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\multicolumn\{\d+\}\{[^}]*\}\{([^{}]*)\}", r"\1", s)
    for cmd in ("\\textbf", "\\textit", "\\emph", "\\texttt", "\\textsuperscript",
                "\\noindent", "\\centering", "\\small", "\\footnotesize", "\\url",
                "\\href", "\\text", "\\mathrm"):
        s = s.replace(cmd, "")
    s = re.sub(r"\\vspace\{[^}]*\}", "", s)
    s = re.sub(r"\\setlength\{[^}]*\}\{[^}]*\}", "", s)
    s = re.sub(r"\\begin\{[^}]*\}\[?[^\]]*\]?", "", s)
    s = re.sub(r"\\end\{[^}]*\}", "", s)
    s = s.replace("---", "—").replace("--", "–").replace("``", "“").replace("''", "”")
    for _k, _v in (("\\,", " "), ("\\;", " "), ("\\ ", " "), ("\\quad", " "),
                   ("\\qquad", " "), ("\\%", "%"), ("\\&", "&"), ("\\_", "_")):
        s = s.replace(_k, _v)
    # elsarticle-harv 的 .bbl 用 \bibinfo{域}{值} + \DOIprefix\doi{...}
    for _ in range(3):
        s = re.sub(r"\\bibinfo\{[^}]*\}\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}", r"\1", s)
        s = re.sub(r"\\bibfield\{[^}]*\}\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}", r"\1", s)
    s = s.replace("\\DOIprefix", "doi: ").replace("\\newblock", " ")
    s = re.sub(r"\\doi\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\[a-zA-Z]+\*?", "", s)
    s = s.replace("{", "").replace("}", "")
    s = s.replace("~", " ").replace(r"\%", "%").replace(r"\&", "&")
    s = re.sub(r"\s+", " ", s).strip()
    _bad = re.findall(r"\\[a-zA-Z]*|\$|\{|\}", s)
    if _bad:
        LEFT.append((s[:90], sorted(set(_bad))))
    return s


# ---------------- 版式（与 build_word_cn.py 完全一致） ----------------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def set_cn(run):
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def h1(text):
    return heading_cn(doc, text, 15)


def h2(text):
    return heading_cn(doc, text, 12.5)


def para(text, size=12, bold=False, align=None, indent=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    set_cn(r)
    if align is not None:
        p.alignment = align
    if indent and align is None:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    return p


def fig(img, caption, w=14.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(FIG, img), width=Cm(w))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(caption)
    r.font.size = Pt(9.5)
    r.bold = True
    set_cn(r)
    return cp


def _vcenter_content(t):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def table(headers, rows, caption, note=None):
    para(caption, 10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        c.paragraphs[0].add_run(htxt).bold = True
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].runs[0].font.size = Pt(9)
        set_cn(c.paragraphs[0].runs[0])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            if i < len(cells):
                r = cells[i].paragraphs[0].add_run(str(v))
                r.font.size = Pt(9)
                set_cn(r)
    make_three_line(t)
    _vcenter_content(t)
    if note:
        para(note, 9, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    return t


def eq(name, num):
    b = {"weight": EQ3, "norm": EQNORM_EN, "esi": EQ4, "gm": EQ5, "q": EQ6}[name]
    return add_equation(doc, b(), num)


EQ_OF = {"eq:weight": ("weight", 1), "eq:norm": ("norm", 2), "eq:esi": ("esi", 3),
         "eq:gm": ("gm", 4), "eq:q": ("q", 5)}

# ---------------- 题头 ----------------
title = re.search(r"\\title\{(.*?)\}\s*\n", TEX, re.S).group(1)
author = re.search(r"\\author\[[^\]]*\]\{(.*)$", TEX, re.M).group(1)
email = re.search(r"\\ead\{(.*?)\}", TEX).group(1)
affil = re.search(r"\\affiliation\[[^\]]*\]\{organization=\{(.*?)\}\}", TEX, re.S).group(1)
abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", TEX, re.S).group(1)
kws = re.search(r"\\begin\{keyword\}(.*?)\\end\{keyword\}", TEX, re.S).group(1)

para(tex2txt(title), 17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para(tex2txt(re.sub(r"\\corref\{[^}]*\}", "", author)) + " (Corresponding author)",
     11, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para(tex2txt(affil), 10, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para("Corresponding author: " + email, 10, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
doc.add_paragraph()
h1("Abstract")
para(tex2txt(abstract))
para("Keywords: " + "; ".join(x for x in (tex2txt(k).strip() for k in kws.split("\\sep")) if x), 10,
     bold=True, indent=False)
doc.add_paragraph()

# ---------------- 正文扫描 ----------------
lines = TEX.split("\n")
start = next(i for i, l in enumerate(lines) if l.strip().startswith("\\end{frontmatter}")) + 1
i = start
sec_n = sub_n = fig_n = tab_n = 0
hl_items = []
buf = []
stats = {"para": 0, "fig": 0, "tab": 0, "eq": 0}


def flush():
    global buf
    if buf:
        t = tex2txt(" ".join(buf))
        if t:
            para(t)
            stats["para"] += 1
        buf = []


def block_end(j, env):
    while j < len(lines) and not lines[j].strip().startswith("\\end{%s}" % env):
        j += 1
    return j


while i < len(lines):
    ln = lines[i].strip()
    if not ln or ln.startswith("%"):
        flush()
        i += 1
        continue
    m = re.match(r"\\section\{(.+)\}$", ln)
    if m:
        flush()
        sec_n += 1
        sub_n = 0
        h1("%d  %s" % (sec_n, tex2txt(m.group(1))))
        i += 1
        continue
    m = re.match(r"\\subsection\{(.+)\}$", ln)
    if m:
        flush()
        sub_n += 1
        h2("%d.%d  %s" % (sec_n, sub_n, tex2txt(m.group(1))))
        i += 1
        continue
    m = re.match(r"\\section\*\{(.+)\}$", ln)
    if m:
        flush()
        _name = tex2txt(m.group(1))
        if _name in BACK_SECTIONS:        # 尾段统一输出，避免正文遍历重复一遍
            i += 1
            while i < len(lines) and not re.match(
                    r"\\(section\*|bibliographystyle|end\{document\})", lines[i].strip()):
                i += 1
            continue
        h1(_name)
        i += 1
        continue
    if ln.startswith("\\begin{table}"):
        flush()
        j = block_end(i, "table")
        body = "\n".join(lines[i:j + 1])
        cap = re.search(r"\\caption\{(.*?)\}\s*\n", body, re.S)
        tab_n += 1
        rows_all = re.findall(r"^[^%\n]*?&[^\n]*$", re.sub(r"\\midrule", "@@MID@@",
                                                            re.sub(r"\\toprule", "@@TOP@@",
                                                                   re.sub(r"\\bottomrule", "@@BOT@@", body))),
                              re.M)
        raw = re.sub(r"\\bottomrule", "\n@@BOT@@", re.sub(r"\\midrule", "\n@@MID@@",
                      re.sub(r"\\toprule", "\n@@TOP@@", body)))
        seg = raw.split("@@MID@@")
        pre = seg[0].split("@@TOP@@")[-1] if "@@TOP@@" in seg[0] else seg[0]
        post = seg[1].split("@@BOT@@")[0] if len(seg) > 1 else ""
        def cells(txt):
            out = []
            for line in txt.split("\n"):
                s = line.strip().rstrip("\\").strip()
                if "&" in s:
                    out.append([tex2txt(c) for c in s.split("&")])
            return out
        hdr = cells(pre)
        bodyrows = [r for r in cells(post) if any(x for x in r)]
        note = None
        mn = re.search(r"\{\\footnotesize(.*?)\}", body, re.S)
        if mn:
            note = tex2txt(mn.group(1))
        table(hdr[0] if hdr else [], bodyrows, "Table %d.  %s" % (tab_n, tex2txt(cap.group(1))), note)
        stats["tab"] += 1
        i = j + 1
        continue
    if ln.startswith("\\begin{figure}"):
        flush()
        j = block_end(i, "figure")
        body = "\n".join(lines[i:j + 1])
        img = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}", body)
        w = re.search(r"width=([\d.]+)\\linewidth", body)
        cap = re.search(r"\\caption\{(.*?)\}\s*\n", body, re.S)
        stem = os.path.splitext(os.path.basename(img.group(1)))[0] + ".png"
        fig_n += 1
        frac = float(w.group(1)) if w else 0.9
        fig(stem, "Figure %d.  %s" % (fig_n, tex2txt(cap.group(1))), w=round(frac * 16.0, 1))
        stats["fig"] += 1
        i = j + 1
        continue
    if ln.startswith("\\begin{equation}"):
        flush()
        j = block_end(i, "equation")
        body = "\n".join(lines[i:j + 1])
        lab = re.search(r"\\label\{([^}]+)\}", body)
        nm, num = EQ_OF.get(lab.group(1), (None, None))
        if nm:
            eq(nm, num)
            stats["eq"] += 1
        i = j + 1
        continue
    if ln.startswith("\\begin{enumerate}"):
        flush()
        j = block_end(i, "enumerate")
        body = "\n".join(lines[i:j + 1])
        items = [tex2txt(x) for x in re.findall(r"\\item\s*(.*?)(?=\\item|\\end\{enumerate\})", body, re.S)]
        items = [x for x in items if x]
        if items:
            para(" ".join("(%d) %s" % (k, t) for k, t in enumerate(items, 1)))
            stats["para"] += 1
        i = j + 1
        continue
    if ln.startswith("\\begin{itemize}"):
        flush()
        j = block_end(i, "itemize")
        body = "\n".join(lines[i:j + 1])
        items = [tex2txt(x) for x in re.findall(r"\\item\s*(.*?)(?=\\item|\\end\{itemize\})", body, re.S)]
        hl_items.extend([x for x in items if x])
        i = j + 1
        continue
    if ln.startswith(("\\bibliographystyle", "\\bibliography", "\\end{document}",
                      "\\label{", "\\vspace", "\\setlength")):
        i += 1
        continue
    buf.append(ln)
    i += 1
flush()

# ---------------- 尾段（镜像 CN Word 顺序） ----------------
doc.add_paragraph()
if hl_items:                      # 正文已无 Highlights 段（Elsevier 要求独立文件）
    h1("Highlights")
    for t in hl_items:
        para("- " + t, 10, indent=False)
doc.add_paragraph()
for name in BACK_SECTIONS:
    m = re.search(r"\\section\*\{%s\}(.*?)(?=\\section\*|\\bibliographystyle|\\end\{document\})"
                  % re.escape(name), TEX, re.S)
    if m:
        h1(name)
        para(tex2txt(m.group(1)))
        stats["para"] += 1
doc.add_paragraph()
h1("References")
refs = []
for k in BBL_KEYS:
    m = re.search(r"\\bibitem(?:\[[^\]]*\])?\{%s\}(.*?)(?=\\bibitem|\\end\{thebibliography\})"
                  % re.escape(k), BBL, re.S)
    if m:
        body = re.sub(r"(?<!\\)%[^\n]*", "", m.group(1))   # 去 .bbl 的 %Type = ... 注释行
        refs.append(tex2txt(body))
for r in refs:                                # 作者-年制（Harvard）：参考文献无序号
    para(r, 9.5, indent=False)

doc.save(OUT)
print("OK:", OUT)
print("段落 %d | 图 %d | 表 %d | 公式 %d | 参考文献 %d | Highlights %d"
      % (stats["para"], stats["fig"], stats["tab"], stats["eq"], len(refs), len(hl_items)))
if LEFT:
    print("[警告] 仍有 %d 处 LaTeX 残留：" % len(LEFT))
    for _s, _b in LEFT[:8]:
        print("   ", _b, _s)
else:
    print("LaTeX 残留检查：干净")
unknown = sorted({k.strip() for m in re.findall(r"\\cite[pt]?\{([^}]*)\}", TEX)
                  for k in m.split(",") if k.strip() and k.strip() not in BIB})
print("未映射引用键:", unknown if unknown else "无")
sys.exit(1 if (LEFT or unknown) else 0)